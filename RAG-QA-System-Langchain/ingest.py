import os
import shutil

from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document

from langchain_text_splitters import (
    RecursiveCharacterTextSplitter,
)

from langchain_google_genai import (
    GoogleGenerativeAIEmbeddings,
)

from langchain_chroma import Chroma

DATA_DIR = "data"
INDEX_DIR = "index"

COLLECTION_NAME = "qa_collection"

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".docx",
}

load_dotenv()

api_key = os.getenv("GOOGLE_API_KEY")
if not api_key:
    raise ValueError("Please configure API key")

print("API key configured successfully")

# ============================================================
# Extract TXT
# ============================================================

def extract_txt(file_path):

    print(f"Reading TXT: {file_path}")

    try:

        with open(
            file_path,
            "r",
            encoding="utf-8",
        ) as file:

            return file.read()

    except UnicodeDecodeError:

        # Fallback for files with another encoding

        with open(
            file_path,
            "r",
            encoding="latin-1",
        ) as file:

            return file.read()


# ============================================================
# Extract PDF
# ============================================================

def extract_pdf(file_path):

    print(f"Reading PDF: {file_path}")

    reader = PdfReader(file_path)

    text = ""

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):

        page_text = page.extract_text() or ""

        print(
            f"    Page {page_number}: "
            f"{len(page_text)} characters"
        )

        text += page_text + "\n"

    return text


# ============================================================
# Extract DOCX
# ============================================================

def extract_docx(file_path):

    print(f"Reading DOCX: {file_path}")

    document = DocxDocument(file_path)

    text_parts = []

    # --------------------------------------------------------
    # Paragraphs
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text_parts.append(
                paragraph.text
            )

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:

                    row_text.append(
                        cell_text
                    )

            if row_text:

                text_parts.append(
                    " | ".join(row_text)
                )

    return "\n".join(text_parts)


# ============================================================
# Extract document based on extension
# ============================================================

def extract_file(file_path):

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        return extract_pdf(file_path)

    elif extension == ".txt":

        return extract_txt(file_path)

    elif extension == ".docx":

        return extract_docx(file_path)

    else:

        print(
            f"Skipping unsupported file: "
            f"{file_path}"
        )

        return ""


# ============================================================
# Find all files recursively
# ============================================================

def find_files():

    files = []

    for root, directories, filenames in os.walk(
        DATA_DIR
    ):

        for filename in filenames:

            file_path = os.path.join(
                root,
                filename,
            )

            extension = os.path.splitext(
                filename
            )[1].lower()

            if extension in SUPPORTED_EXTENSIONS:

                files.append(file_path)

    return files

# ============================================================
# Main ingestion
# ============================================================

print("\n" + "=" * 70)
print("DOCUMENT INGESTION")
print("=" * 70)

# ------------------------------------------------------------
# Check data directory
# ------------------------------------------------------------

if not os.path.exists(DATA_DIR):

    raise FileNotFoundError(
        f"Data directory not found: {DATA_DIR}"
    )

# ------------------------------------------------------------
# Find files
# ------------------------------------------------------------

files = find_files()

print(
    f"\nFound {len(files)} supported file(s)."
)

if not files:

    raise ValueError(
        "No supported files found in data/."
    )

# ============================================================
# Read all files
# ============================================================

documents = []

successful_files = 0
failed_files = 0


for file_path in files:

    print("\n" + "-" * 70)

    try:

        text = extract_file(file_path)

        text = text.strip()

        if not text:

            print(
                "WARNING: No text extracted."
            )

            failed_files += 1

            continue


        # ----------------------------------------------------
        # Relative path for metadata
        # ----------------------------------------------------

        relative_path = os.path.relpath(
            file_path,
            DATA_DIR,
        )


        # ----------------------------------------------------
        # Create LangChain Document
        # ----------------------------------------------------

        document = Document(
            page_content=text,
            metadata={
                "file_name": os.path.basename(
                    file_path
                ),
                "file_path": relative_path,
                "file_type": os.path.splitext(
                    file_path
                )[1].lower(),
            },
        )


        documents.append(
            document
        )

        successful_files += 1


        print(
            f"Extracted characters: {len(text)}"
        )

        print(
            f"Source: {relative_path}"
        )


    except Exception as e:

        failed_files += 1

        print(
            f"ERROR reading {file_path}: {e}"
        )


# ============================================================
# Summary
# ============================================================

print("\n" + "=" * 70)
print("DOCUMENT SUMMARY")
print("=" * 70)

print(
    "Files found:       ",
    len(files)
)

print(
    "Successfully read: ",
    successful_files
)

print(
    "Failed:             ",
    failed_files
)

print(
    "LangChain docs:     ",
    len(documents)
)


if not documents:

    raise ValueError(
        "No documents could be extracted."
    )


# ============================================================
# Split documents
# ============================================================

print("\n" + "=" * 70)
print("SPLITTING DOCUMENTS")
print("=" * 70)


splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=100,
)


chunks = splitter.split_documents(
    documents
)


print(
    f"Created {len(chunks)} chunks."
)

# ============================================================
# Show sample chunks
# ============================================================

print("\nSample chunks:")

for i, chunk in enumerate(
    chunks[:5],
    start=1,
):

    print("\n" + "-" * 70)

    print(f"CHUNK {i}")

    print(
        "Source:",
        chunk.metadata.get(
            "file_name",
            "unknown",
        ),
    )

    print(
        "Characters:",
        len(chunk.page_content),
    )

    print(
        chunk.page_content[:500]
    )

# ============================================================
# Create embedding model
# ============================================================

print("\n" + "=" * 70)
print("INITIALIZING EMBEDDING MODEL")
print("=" * 70)


embed_model = GoogleGenerativeAIEmbeddings(
    model="gemini-embedding-001",
)


print(
    "Embedding model initialized."
)

# ============================================================
# Remove old index
# ============================================================

if os.path.exists(INDEX_DIR):

    print(
        f"\nRemoving existing index: "
        f"{INDEX_DIR}"
    )

    shutil.rmtree(
        INDEX_DIR
    )


# ============================================================
# Create Chroma vector database
# ============================================================

print("\n" + "=" * 70)
print("CREATING VECTOR DATABASE")
print("=" * 70)


vectorstore = Chroma.from_documents(
    documents=chunks,
    embedding=embed_model,
    persist_directory=INDEX_DIR,
    collection_name=COLLECTION_NAME,
)


print(
    "Chroma vector database created."
)

# ============================================================
# Final summary
# ============================================================

print("\n" + "=" * 70)
print("INGESTION COMPLETED")
print("=" * 70)

print(
    f"Files processed: {successful_files}"
)

print(
    f"Files failed:    {failed_files}"
)

print(
    f"Chunks created:  {len(chunks)}"
)

print(
    f"Index location:  {INDEX_DIR}/"
)

print(
    "Vector database:  ChromaDB"
)

print("=" * 70)